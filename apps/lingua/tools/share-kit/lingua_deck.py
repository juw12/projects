"""Lingua slide decks: the installation & usage guide, and per-release update decks.

Used by build_deck.py (first share) and notifier/lingua_update_notifier.py (update e-mails).
"""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
IMG = HERE / "img"
SHOTS = {
    "main": IMG / "main.png",
    "text": IMG / "text.png",
    "alts": IMG / "alts.png",
    "picture": IMG / "picture.png",
    "phone": IMG / "phone.png",
}

NAVY = RGBColor(0x0F, 0x17, 0x2A)
ACCENT = RGBColor(0x0E, 0x86, 0xD4)
LIGHT = RGBColor(0x7D, 0xD3, 0xFC)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PANEL = RGBColor(0xF1, 0xF5, 0xF9)
BORDER = RGBColor(0xCB, 0xD5, 0xE1)
WARN = RGBColor(0xB4, 0x53, 0x09)
GREEN = RGBColor(0x15, 0x80, 0x3D)
YELLOW = RGBColor(0xFD, 0xE6, 0x8A)

SW = Inches(13.333)
SH = Inches(7.5)
FOOTER = {"text": "Lingua"}
# Contact details shown on the last slide and title slide. Override locally (e.g. from the
# notifier config) - the public defaults contain no personal e-mail address.
CONTACT = {
    "line": "Juwen Lie  ·  github.com/juw12",
    "note": "Prepared by Juwen Lie",
}


# ----------------------------------------------------------------- helpers
def new_presentation():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_text(slide, left, top, width, height, paragraphs, size=18, color=INK,
             bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Segoe UI"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    if isinstance(paragraphs, str):
        paragraphs = [paragraphs]
    for index, item in enumerate(paragraphs):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(6)
        runs = item if isinstance(item, list) else [(item, {})]
        for run_text, style in runs:
            run = paragraph.add_run()
            run.text = run_text
            run.font.name = style.get("font", font)
            run.font.size = Pt(style.get("size", size))
            run.font.bold = style.get("bold", bold)
            run.font.color.rgb = style.get("color", color)
    return box


def background(slide, color=NAVY):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    rect.fill.solid()
    rect.fill.fore_color.rgb = color
    rect.line.fill.background()


def header(slide, title, section):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(1.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    add_text(slide, Inches(0.5), Inches(0.12), Inches(9.5), Inches(0.8), title,
             size=28, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(9.6), Inches(0.12), Inches(3.3), Inches(0.8), section,
             size=14, color=LIGHT, bold=True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(0.5), Inches(7.05), Inches(12.3), Inches(0.35),
             FOOTER["text"], size=10, color=MUTED)


def picture_fit(slide, path, left, top, max_w, max_h, border=True):
    img = Image.open(path)
    ratio = min(max_w / img.width, max_h / img.height)
    width = int(img.width * ratio)
    height = int(img.height * ratio)
    x = left + (max_w - width) // 2
    y = top + (max_h - height) // 2
    pic = slide.shapes.add_picture(str(path), x, y, width, height)
    if border:
        pic.line.color.rgb = BORDER
        pic.line.width = Pt(1)
    return pic


def step_badge(slide, left, top, number, size=Inches(0.55)):
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    circle.fill.solid()
    circle.fill.fore_color.rgb = ACCENT
    circle.line.fill.background()
    tf = circle.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = tf.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = str(number)
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = WHITE
    run.font.name = "Segoe UI"


def steps(slide, items, left, top, width, gap=Inches(1.0), title_size=19, body_size=14):
    for index, (title, body) in enumerate(items):
        y = top + gap * index
        step_badge(slide, left, y, index + 1)
        paragraphs = [[(title, {"bold": True, "size": title_size})]]
        if body:
            paragraphs.append([(body, {"size": body_size, "color": MUTED})])
        add_text(slide, left + Inches(0.75), y - Inches(0.05), width - Inches(0.75), gap, paragraphs)


def callout(slide, left, top, width, height, title, body, color=ACCENT):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = color
    shape.line.width = Pt(1.5)
    paragraphs = [[(title, {"bold": True, "size": 16, "color": color})]]
    paragraphs += [[(line, {"size": 14})] for line in body]
    add_text(slide, left + Inches(0.2), top + Inches(0.12), width - Inches(0.4), height - Inches(0.2), paragraphs)


def bullets(slide, left, top, width, height, items, size=17):
    paragraphs = []
    for item in items:
        head, tail = item if isinstance(item, tuple) else ("", item)
        paragraphs.append([
            ("•  ", {"color": ACCENT, "bold": True, "size": size}),
            (head, {"bold": True, "size": size}),
            (tail, {"size": size}),
        ])
    return add_text(slide, left, top, width, height, paragraphs)


def table_slide(slide, headers, rows, col_widths, top=Inches(1.4), font_size=14):
    total_height = Inches(0.5) * (len(rows) + 1)
    table = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(0.6), top,
                                   sum(col_widths, Inches(0)), total_height).table
    for col, width in enumerate(col_widths):
        table.columns[col].width = width
    for col, label in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = label
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.color.rgb = WHITE
        run.font.bold = True
        run.font.size = Pt(font_size + 1)
        run.font.name = "Segoe UI"
    for row_index, row in enumerate(rows, start=1):
        highlight = isinstance(row, dict) and row.get("highlight")
        values = row["values"] if isinstance(row, dict) else row
        for col, value in enumerate(values):
            cell = table.cell(row_index, col)
            cell.text = value
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if highlight:
                cell.fill.fore_color.rgb = RGBColor(0xDB, 0xEE, 0xFE)
            else:
                cell.fill.fore_color.rgb = PANEL if row_index % 2 else WHITE
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt(font_size)
            run.font.bold = col == 0
            run.font.color.rgb = INK
            run.font.name = "Segoe UI"
    return table


def title_slide(prs, title, subtitle, tagline, note):
    s = blank(prs)
    background(s)
    logo = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.9), Inches(1.9), Inches(1.2), Inches(1.2))
    logo.fill.solid()
    logo.fill.fore_color.rgb = ACCENT
    logo.line.fill.background()
    logo.text_frame.text = "L"
    logo_paragraph = logo.text_frame.paragraphs[0]
    logo_paragraph.alignment = PP_ALIGN.CENTER
    logo_paragraph.runs[0].font.size = Pt(48)
    logo_paragraph.runs[0].font.bold = True
    logo_paragraph.runs[0].font.color.rgb = WHITE
    add_text(s, Inches(2.4), Inches(1.75), Inches(7), Inches(1.0), title, size=54, color=WHITE, bold=True)
    add_text(s, Inches(2.4), Inches(2.75), Inches(7), Inches(0.6), subtitle, size=22, color=LIGHT)
    add_text(s, Inches(0.9), Inches(4.0), Inches(7.8), Inches(1.8), tagline, size=20, color=WHITE)
    add_text(s, Inches(0.9), Inches(6.4), Inches(7.8), Inches(0.5), note, size=14, color=MUTED)
    picture_fit(s, SHOTS["main"], Inches(9.0), Inches(1.6), Inches(3.9), Inches(4.4))
    return s


def contact_slide(prs, release_tag):
    s = blank(prs)
    background(s)
    add_text(s, Inches(0.9), Inches(2.3), Inches(11.5), Inches(1), "Questions or feedback?",
             size=44, color=WHITE, bold=True)
    add_text(s, Inches(0.9), Inches(3.5), Inches(11.5), Inches(2), [
        CONTACT["line"],
        f"Release notes: github.com/juw12/projects/releases  ({release_tag})",
    ], size=22, color=LIGHT)


def section_divider(prs, title, subtitle):
    s = blank(prs)
    background(s, ACCENT)
    add_text(s, Inches(0.9), Inches(2.6), Inches(11.5), Inches(1.2), title, size=44, color=WHITE, bold=True)
    add_text(s, Inches(0.9), Inches(3.8), Inches(11.5), Inches(1.0), subtitle, size=22, color=WHITE)


# ------------------------------------------------------------ guide slides
def add_guide_slides(prs, version):
    """Installation + usage guide (everything after the title slide)."""
    release_tag = f"lingua-v{version}"

    s = blank(prs)
    header(s, "What Lingua does", "OVERVIEW")
    cards = [
        ("Live captions", "Reads the Microsoft Teams live-caption strip and shows it translated, as people talk."),
        ("Text", "Paste or type text → translation with tone, context, glossary, alternatives & word notes."),
        ("Read picture", "Photo, screenshot or screen area → the text in it (OCR), optionally translated."),
        ("Phone", "Scan a QR code to use your phone camera; the API key stays on the laptop."),
    ]
    card_w = Inches(2.9)
    for index, (title, body) in enumerate(cards):
        x = Inches(0.5) + index * (card_w + Inches(0.23))
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.5), card_w, Inches(2.4))
        card.adjustments[0] = 0.06
        card.fill.solid()
        card.fill.fore_color.rgb = PANEL
        card.line.color.rgb = BORDER
        add_text(s, x + Inches(0.2), Inches(1.65), card_w - Inches(0.4), Inches(2.2),
                 [[(title, {"bold": True, "size": 21, "color": ACCENT})], [(body, {"size": 15})]])
    bullets(s, Inches(0.6), Inches(4.3), Inches(12), Inches(2.6), [
        ("Languages: ", "Auto-detect, English, 繁體中文, Bahasa Indonesia (16 targets for pictures)."),
        ("Engines: ", "company AI models first → local AI on your PC (Ollama, optional) → offline packs."),
        ("Privacy: ", "online services (Google/MyMemory) are OFF by default; nothing typed in Text is saved."),
        ("Updates: ", "built-in — Lingua checks GitHub on start and every 4 h; one click to update."),
    ])

    s = blank(prs)
    header(s, "Before you start", "INSTALLATION")
    bullets(s, Inches(0.6), Inches(1.4), Inches(7.5), Inches(5.5), [
        ("Windows 10 / 11 laptop. ", "No admin rights and no Python needed."),
        ("Office network or VPN. ", "The company AI gateway is only reachable from inside."),
        ("lingua.env ", "— gateway settings file, from the first Lingua e-mail (contains NO API key)."),
        ("Your own MaaS API key. ", "If you already use OpenCode with the company gateway, Lingua finds it automatically. Otherwise see 'Connect your API key'."),
        ("Microsoft Teams ", "with Live captions, for caption translation."),
        ("Optional: Ollama ", "(ollama.com) + 'ollama pull qwen3:8b' — local AI when company models are down."),
    ], size=17)
    callout(s, Inches(8.5), Inches(1.5), Inches(4.3), Inches(2.5), "Download",
            ["Lingua-windows.zip on GitHub:", "github.com/juw12/projects/releases", f"→ {release_tag}",
             "(direct link is in the e-mail)"])
    callout(s, Inches(8.5), Inches(4.3), Inches(4.3), Inches(2.3), "Keep your key private",
            ["Never forward a lingua.env that has", "your MAAS_API_KEY filled in —",
             "anyone with the file can use it."], color=WARN)

    s = blank(prs)
    header(s, "Install in 4 steps", "INSTALLATION")
    steps(s, [
        ("Download Lingua-windows.zip", f"Click the link in the e-mail (GitHub release {release_tag})."),
        ("Unzip it", "Right-click the zip → Extract All…  You get a folder called Lingua."),
        ("Put lingua.env inside that folder", "Save the e-mail attachment next to Lingua.exe. Keep the name exactly 'lingua.env'."),
        ("Double-click 'Install Lingua.cmd'", "Installs to %LOCALAPPDATA%\\Programs\\Lingua with Start-menu + Desktop shortcuts and an entry in Settings → Apps. Then you can delete the zip."),
    ], Inches(0.6), Inches(1.45), Inches(7.7), gap=Inches(1.3))
    callout(s, Inches(8.6), Inches(1.5), Inches(4.2), Inches(2.4), "SmartScreen warning?",
            ["The app is not code-signed, so you may", "see 'Windows protected your PC'.",
             "Click  More info  →  Run anyway."], color=WARN)
    callout(s, Inches(8.6), Inches(4.2), Inches(4.2), Inches(2.3), "Firewall prompt?",
            ["When you use the Phone button,", "Windows may ask to allow Lingua.",
             "Allow on Private networks only."])

    s = blank(prs)
    header(s, "What the folder should look like", "INSTALLATION")
    panel = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.5), Inches(6.9), Inches(4.7))
    panel.adjustments[0] = 0.04
    panel.fill.solid()
    panel.fill.fore_color.rgb = NAVY
    panel.line.fill.background()
    mono = {"font": "Consolas", "size": 17, "color": WHITE}
    add_text(s, Inches(0.9), Inches(1.7), Inches(6.5), Inches(4.4), [
        [("Lingua\\", dict(mono, color=LIGHT, bold=True, size=18))],
        [("   _internal\\", mono)],
        [("   Install Lingua.cmd   ← double-click", mono)],
        [("   Lingua.exe", mono)],
        [("   lingua.env           ← from the e-mail", dict(mono, color=YELLOW, bold=True))],
        [("   lingua.env.example", mono)],
        [("   README.txt", mono)],
        [("   THIRD_PARTY_NOTICES.txt", mono)],
    ])
    bullets(s, Inches(7.9), Inches(1.5), Inches(5.0), Inches(5), [
        ("Portable option: ", "skip the installer and just double-click Lingua.exe in this folder."),
        ("Alternative location: ", "lingua.env also works in %APPDATA%\\lingua\\lingua.env."),
        ("Uninstall: ", "Settings → Apps → Lingua → Uninstall."),
        ("Updates keep ", "your settings, history and lingua.env."),
    ], size=17)

    s = blank(prs)
    header(s, "Connect your API key", "INSTALLATION")
    add_text(s, Inches(0.6), Inches(1.3), Inches(12), Inches(0.6),
             "Lingua looks for your MaaS key in this order — use whichever is easiest:", size=18, color=MUTED)
    steps(s, [
        ("OpenCode already set up?  Nothing to do.", "Lingua reads the key from %USERPROFILE%\\.config\\opencode\\opencode.json automatically."),
        ("Or set a Windows environment variable", "Start → 'Edit environment variables for your account' → New → Name: MAAS_API_KEY, Value: your key. Restart Lingua."),
        ("Or paste it into lingua.env", "Open lingua.env in Notepad, fill in the line  MAAS_API_KEY=your-key  and save. Keep this file private."),
    ], Inches(0.6), Inches(2.05), Inches(12), gap=Inches(1.3))
    callout(s, Inches(0.6), Inches(5.95), Inches(12.1), Inches(0.95), "No key yet?",
            ["Text translation still works with local AI (Ollama) or offline language packs — slower, lower quality."])

    s = blank(prs)
    header(s, "The main window", "USAGE")
    picture_fit(s, SHOTS["main"], Inches(0.5), Inches(1.35), Inches(6.4), Inches(5.6))
    bullets(s, Inches(7.2), Inches(1.35), Inches(5.8), Inches(5.6), [
        ("Start ", "— begin / pause live caption translation."),
        ("Region ", "— drag a box around the Teams caption strip (once)."),
        ("Read picture ", "— OCR a photo, screenshot or screen area."),
        ("Text ", "— text translation desk (since 2.3, next to Read picture)."),
        ("Phone ", "— QR code to use your phone camera."),
        ("From / To ", "— source (Auto-detect) and target language."),
        ("More ▾ ", "— model, opacity, font size, click-through."),
        ("ⓘ  /  ☀ ", "— settings, offline packs, updates · light/dark."),
    ], size=16)

    s = blank(prs)
    header(s, "Translate a Teams meeting live", "USAGE")
    steps(s, [
        ("Turn on Teams live captions", "In the meeting: … More → Language and speech → Turn on live captions."),
        ("Pick the caption region (first time only)", "Click Region and drag a rectangle tightly around the caption strip. Lingua remembers it."),
        ("Choose languages", "From: Auto (detect) is fine.   To: your language (e.g. English or 繁體中文)."),
        ("Press Start", "Translated lines appear in Lingua, newest at the bottom. Scroll up to read back."),
    ], Inches(0.6), Inches(1.45), Inches(7.7), gap=Inches(1.3))
    callout(s, Inches(8.6), Inches(1.5), Inches(4.2), Inches(5.0), "Tips",
            ["• Keep the Teams captions visible —", "   Lingua reads the pixels on screen.",
             "• Don't cover them with Lingua; use", "   More ▾ → opacity + Click-through.",
             "• Moved / resized Teams?", "   Click Region again.",
             "• Clear wipes the transcript.",
             "• Status bar shows the speed."])

    s = blank(prs)
    header(s, "Translate text  (Text button)", "USAGE")
    picture_fit(s, SHOTS["text"], Inches(0.5), Inches(1.3), Inches(7.6), Inches(5.65))
    bullets(s, Inches(8.4), Inches(1.35), Inches(4.6), Inches(5.6), [
        ("Tone: ", "Natural, Professional, Friendly, Formal, Literal."),
        ("Ctrl+Enter ", "translates, Esc stops."),
        ("Context ", "— who it is for / what it is about."),
        ("Glossary ", "— one per line:  PCB layout = 電路板佈局"),
        ("Engine: ", "'Best available' picks automatically."),
        ("Paste · Copy · Save .txt · Clear", ""),
        ("Warnings ", "if numbers, links or glossary terms changed."),
    ], size=16)

    s = blank(prs)
    header(s, "Alternatives & word-choice notes", "USAGE")
    picture_fit(s, SHOTS["alts"], Inches(0.5), Inches(1.3), Inches(7.6), Inches(5.65))
    bullets(s, Inches(8.4), Inches(1.35), Inches(4.6), Inches(5.6), [
        ("Click 'Alternatives & notes' ", "after translating."),
        ("Two other phrasings ", "plus English notes on word choice (e.g. what sounds more natural in Taiwan)."),
        ("Takes up to ~1 min ", "on local AI; faster on company models."),
        ("Nothing typed here is saved.", ""),
    ], size=16)

    s = blank(prs)
    header(s, "Read a picture (OCR)", "USAGE")
    picture_fit(s, SHOTS["picture"], Inches(0.5), Inches(1.3), Inches(7.8), Inches(5.65))
    bullets(s, Inches(8.6), Inches(1.35), Inches(4.4), Inches(5.0), [
        ("Open ", "a photo / screenshot (multi-select = batch), Paste (Ctrl+V) or Grab screen area."),
        ("Rotate / crop ", "if needed, then Read."),
        ("Text tab ", "— words, letters, line count."),
        ("Translation tab ", "— pick a language or tick Translate automatically."),
        ("Listen · Copy · Save .txt · History", ""),
    ], size=16)
    add_text(s, Inches(8.6), Inches(6.3), Inches(4.4), Inches(0.6),
             "Screenshot: phone / web version. The desktop window has the same Text / Translation tabs.",
             size=11, color=MUTED)

    s = blank(prs)
    header(s, "Use your phone camera  (Phone button)", "USAGE")
    picture_fit(s, SHOTS["phone"], Inches(0.6), Inches(1.3), Inches(4.4), Inches(5.65))
    steps(s, [
        ("Click Phone on the laptop", "A QR code appears. Phone and laptop must be on the same Wi-Fi."),
        ("Scan the QR code with the phone camera", "The Lingua picture page opens in the browser."),
        ("Add to Home Screen", "iPhone: Share → Add to Home Screen.   Android: ⋮ menu → Add to Home screen."),
        ("Take a photo → Read → Translate", "Works while the Phone window is open; Stop sharing ends it."),
    ], Inches(5.5), Inches(1.45), Inches(7.4), gap=Inches(1.3))

    s = blank(prs)
    header(s, "Updates & when the company AI is down", "USAGE")
    callout(s, Inches(0.6), Inches(1.5), Inches(5.9), Inches(3.9), "Updating",
            ["• Lingua checks on start and every 4 hours.",
             "• Click ⟳ Updates (top right) to check now.",
             "• A blue ⬆ Update button + banner appears",
             "   when a new version is out.",
             "• Click Update now: it downloads, verifies,",
             "   restarts and keeps your settings,",
             "   history and lingua.env."])
    callout(s, Inches(6.9), Inches(1.5), Inches(5.9), Inches(3.9), "Fallback order (automatic)",
            ["1. Company AI models (best).",
             "2. Local AI on your PC — install Ollama,",
             "    then run:  ollama pull qwen3:8b",
             "3. Offline language packs (~70 MB / pair):",
             "    ⓘ → Translation → Download packs",
             "    for my languages.",
             "4. Online services — OFF unless you allow",
             "    them (they send text outside the company)."])

    s = blank(prs)
    header(s, "Troubleshooting", "HELP")
    rows = [
        ("'Windows protected your PC'", "More info → Run anyway (the app is not code-signed)."),
        ("Translation unavailable / errors", "Office network or VPN on? lingua.env next to Lingua.exe? API key set?"),
        ("Nothing appears in a Teams call", "Live captions on? Region still over the caption strip? Pressed Start?"),
        ("First translation is slow", "Local AI / offline packs load once (20–35 s); then 2–3 s each."),
        ("Phone can't open the link", "Same Wi-Fi as the laptop; allow Lingua in the firewall (private); keep the Phone window open."),
        ("Odd translation", "Set From to the real language, add Context / Glossary, or try Alternatives."),
        ("Update button does nothing", "Check office network / VPN, then ⓘ → Check for updates; or re-download the zip and run Install Lingua.cmd again."),
    ]
    table_slide(s, ["Problem", "Fix"], rows, [Inches(4.0), Inches(8.1)], font_size=14)


# --------------------------------------------------------- update slides
def add_update_slides(prs, new_version, previous_version, changes, history, download_url):
    """changes: [(version, date, headline, [bullet, ...]), ...] newest first, only versions newer than previous.
    history: [(version, date, headline), ...] newest first, all releases."""
    s = blank(prs)
    header(s, f"How to update to {new_version}", "UPDATE")
    steps(s, [
        ("Open Lingua", "Start menu → Lingua (or the Desktop shortcut)."),
        ("Click ⟳ Updates (top right)", f"It turns into  ⬆ Update to {new_version}  with a blue banner. Lingua also checks by itself on start and every 4 hours."),
        ("Click Update now", "Lingua downloads the new version, verifies it, and restarts. Settings, history and lingua.env are kept."),
        ("Check the version", f"The header now shows v{new_version}."),
    ], Inches(0.6), Inches(1.45), Inches(7.7), gap=Inches(1.3))
    callout(s, Inches(8.6), Inches(1.5), Inches(4.2), Inches(2.6), "Update button not working?",
            ["Download the zip again from", "github.com/juw12/projects/releases",
             f"→ lingua-v{new_version} (link in the e-mail).",
             "Unzip, copy your lingua.env in,", "run Install Lingua.cmd."], color=WARN)
    callout(s, Inches(8.6), Inches(4.4), Inches(4.2), Inches(2.1), "Not installed yet?",
            ["Follow the installation slides", "later in this deck (lingua.env is in", "the first Lingua e-mail)."])

    for version, date, headline, items in changes:
        pages = [items[i:i + 6] for i in range(0, max(len(items), 1), 6)] or [[]]
        for page_index, page in enumerate(pages):
            s = blank(prs)
            suffix = "" if len(pages) == 1 else f"  ({page_index + 1}/{len(pages)})"
            label = "NEW" if version == new_version else "ALSO NEW"
            header(s, f"What's new in {version}{suffix}", label)
            add_text(s, Inches(0.6), Inches(1.3), Inches(12.1), Inches(0.9),
                     [[(headline, {"bold": True, "size": 20, "color": ACCENT})],
                      [(f"Released {date}", {"size": 13, "color": MUTED})]])
            if page:
                bullets(s, Inches(0.6), Inches(2.35), Inches(12.1), Inches(4.6), page, size=17)
            else:
                add_text(s, Inches(0.6), Inches(2.35), Inches(12.1), Inches(1),
                         "Fixes and improvements (see release notes).", size=17)

    s = blank(prs)
    header(s, "Version history", "HISTORY")
    rows = []
    for version, date, headline in history[:10]:
        marker = ""
        if version == new_version:
            marker = "  ← new"
        elif version == previous_version:
            marker = "  ← you were sent"
        rows.append({"values": (version + marker, date, headline), "highlight": version == new_version})
    table_slide(s, ["Version", "Date", "Highlights"], rows,
                [Inches(2.6), Inches(1.7), Inches(7.8)], font_size=13)


# ---------------------------------------------------------------- builders
def build_guide_deck(version, out_path):
    FOOTER["text"] = f"Lingua {version}  ·  Installation & Usage Guide"
    prs = new_presentation()
    title_slide(prs, "Lingua", f"Installation & Usage Guide  ·  Windows · v{version}",
                ["Translate live Teams captions, typed text and pictures —",
                 "English · 繁體中文 · Bahasa Indonesia and more."],
                CONTACT["note"])
    add_guide_slides(prs, version)
    contact_slide(prs, f"lingua-v{version}")
    prs.save(out_path)
    return out_path


def build_update_deck(new_version, previous_version, changes, history, download_url, out_path):
    FOOTER["text"] = f"Lingua {new_version}  ·  Update, Changes & Usage Guide"
    prs = new_presentation()
    title_slide(prs, "Lingua update", f"Version {new_version}  ·  Changes & Usage Guide",
                [changes[0][2] if changes else "A new version of Lingua is available.",
                 f"Updating from {previous_version}? Two clicks: ⟳ Updates → Update now."],
                CONTACT["note"])
    add_update_slides(prs, new_version, previous_version, changes, history, download_url)
    section_divider(prs, "Usage guide", f"Full installation & usage guide for Lingua {new_version}")
    add_guide_slides(prs, new_version)
    contact_slide(prs, f"lingua-v{new_version}")
    prs.save(out_path)
    return out_path
