#!/usr/bin/env python3
"""Build the free Stackup Lite workbook (Excel + Google Sheets compatible).

Build-time only: needs openpyxl (`pip install openpyxl`). The web app and the
tests do not depend on it. Output: apps/stackup/web/downloads/Stackup-Lite.xlsx

Formulas stick to functions that exist in both Excel and Google Sheets
(SUMPRODUCT, SQRT, NORMSDIST, INDEX/MATCH, IF/AND/OR) so the same file works
after "File > Import" in Sheets.

`add_start_sheet` and `add_stack_sheet` are reused by the paid Pro builder.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

APP_URL = "https://juwpegatron.github.io/projects/stackup/"
FIRST_ROW = 12
ROW_COUNT = 20
LAST_ROW = FIRST_ROW + ROW_COUNT - 1

NAVY = "0B1220"
INK = "0F172A"
MUTED = "55627A"
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")
CALC_FILL = PatternFill("solid", fgColor="EEF2F8")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
RESULT_FILL = PatternFill("solid", fgColor="E8F4FD")
THIN = Side(style="thin", color="D6DCEB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WHITE_BOLD = Font(bold=True, color="FFFFFF")
TITLE = Font(bold=True, size=16, color=INK)
SUBTLE = Font(italic=True, size=9, color=MUTED)
BOLD = Font(bold=True, color=INK)

GASKET_EXAMPLE = {
    "name": "IPX7 front-cover gasket compression (example)",
    "units": "mm",
    "lsl": 0.35,
    "usl": 0.65,
    "target_cpk": 1.33,
    "rows": [
        ("Gasket free thickness", "Gasket", 1.6, 0.08, 0.08, 1, "Normal", 3),
        ("Groove depth", "Housing", 1.2, 0.05, 0.05, -1, "Normal", 3),
        ("Compression rib height", "Front cover", 0.3, 0.05, 0.05, 1, "Normal", 3),
        ("Cover stand-off at screw boss", "Front cover", 0.2, 0.05, 0.05, -1, "Normal", 3),
        ("Seal-face flatness", "Housing", 0, 0.05, 0.05, -1, "Uniform", 3),
    ],
}


def define_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def style_range(ws, ref, fill=None, font=None, border=BOX, number_format=None, align=None):
    if ":" not in ref:
        ref = f"{ref}:{ref}"
    for row in ws[ref]:
        for cell in row:
            if fill:
                cell.fill = fill
            if font:
                cell.font = font
            if border:
                cell.border = border
            if number_format:
                cell.number_format = number_format
            if align:
                cell.alignment = align


def add_start_sheet(wb, edition, extra_lines=()):
    ws = wb.active
    ws.title = "Start"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110
    ws["B2"] = f"Stackup {edition} \u2014 tolerance stack-up workbook"
    ws["B2"].font = TITLE
    ws["B3"] = f"Excel + Google Sheets \u00b7 companion to the free web app: {APP_URL}"
    ws["B3"].font = SUBTLE
    lines = [
        ("How to use", True),
        ("1. Go to the Stack sheet. Yellow cells are inputs, grey cells are calculated \u2014 do not type over them.", False),
        ("2. Enter the stack name, units and the spec limits (LSL and/or USL) of the gap you care about. Leave a limit blank if the spec is one-sided.", False),
        ("3. List each dimension in the loop: nominal, plus tolerance, minus tolerance (both as positive numbers).", False),
        ("4. Sensitivity: +1 if the dimension opens the gap, \u22121 if it closes it, 0.5 for a diameter acting as a radius.", False),
        ("5. Distribution: Normal for capable molded / machined features (k = 3 means the tolerance is \u00b13\u03c3, k = 4 means Cp 1.33), Uniform for flatness, float and clearance fits, Triangular for hand assembly.", False),
        ("6. Read the verdict: worst case, RSS (\u00b13\u03c3), Cpk and predicted ppm out of spec. The variance-share column shows which dimension to tighten first.", False),
        ("", False),
        ("Google Sheets", True),
        ("Drive > New > File upload > this .xlsx, then open it with Google Sheets (or File > Import in an open sheet). All formulas use functions both apps support.", False),
        ("", False),
        ("Method", True),
        ("Mean shift for asymmetric tolerances: mean = nominal + (tol+ \u2212 tol\u2212)/2, half-band = (tol+ + tol\u2212)/2.", False),
        ("Worst case = \u03a3 a\u00b7mean \u00b1 \u03a3 |a|\u00b7half. RSS \u03c3 = \u221a\u03a3 (a\u00b7\u03c3\u1d62)\u00b2 with \u03c3\u1d62 = half/k (normal), half/\u221a3 (uniform), half/\u221a6 (triangular).", False),
        ("Cpk = min(USL \u2212 mean, mean \u2212 LSL) / 3\u03c3. Predicted ppm assumes the assembled gap is normal \u2014 non-normal inputs (uniform) usually make real tails lighter; use Monte Carlo to check.", False),
        ("", False),
    ]
    lines += [(text, bold) for text, bold in extra_lines]
    lines += [
        ("Disclaimer & licence", True),
        ("Engineering aid only \u2014 not a substitute for a reviewed tolerance analysis. Example values are illustrative, not from a real product.", False),
        ("Free to use for personal and commercial engineering work. Please do not resell or redistribute the template itself. \u00a9 Juwen Lie.", False),
    ]
    row = 5
    for text, bold in lines:
        cell = ws.cell(row=row, column=2, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = Font(bold=True, size=12, color=INK) if bold else Font(color=INK)
        row += 1
    return ws


def add_stack_sheet(wb, example=GASKET_EXAMPLE, title="Stack"):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    widths = {
        "A": 4, "B": 30, "C": 16, "D": 11, "E": 9, "F": 9, "G": 8, "H": 13, "I": 7,
        "J": 11, "K": 10, "L": 10, "M": 11, "N": 10, "O": 12, "P": 11, "Q": 3, "R": 30, "S": 22,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws["A1"] = "Tolerance stack-up"
    ws["A1"].font = TITLE
    ws["A2"] = f"Stackup \u00b7 yellow = input \u00b7 grey = calculated \u00b7 web app: {APP_URL}"
    ws["A2"].font = SUBTLE

    labels = [
        ("B4", "Stack name", "C4", example["name"]),
        ("B5", "Units", "C5", example["units"]),
        ("B6", "Lower spec limit (LSL) \u2014 blank if none", "C6", example["lsl"]),
        ("B7", "Upper spec limit (USL) \u2014 blank if none", "C7", example["usl"]),
        ("B8", "Target Cpk", "C8", example["target_cpk"]),
    ]
    for label_ref, label, value_ref, value in labels:
        ws[label_ref] = label
        ws[label_ref].font = BOLD
        ws[value_ref] = value
        style_range(ws, value_ref, fill=INPUT_FILL)
    ws.merge_cells("C4:H4")
    style_range(ws, "C4:H4", fill=INPUT_FILL)

    sheet = f"'{title}'"
    define_name(wb, "LSL", f"{sheet}!$C$6")
    define_name(wb, "USL", f"{sheet}!$C$7")
    define_name(wb, "TargetCpk", f"{sheet}!$C$8")

    headers = [
        "#", "Dimension", "Part / source", "Nominal", "Tol +", "Tol \u2212", "Sens.", "Distribution", "k",
        "Mean", "Half band", "\u03c3", "a\u00b7mean", "|a|\u00b7half", "(a\u00b7\u03c3)\u00b2", "Var. share",
    ]
    for index, header in enumerate(headers, start=1):
        cell = ws.cell(row=FIRST_ROW - 1, column=index, value=header)
        cell.fill = HEAD_FILL
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX
    ws.row_dimensions[FIRST_ROW - 1].height = 30
    ws.cell(row=FIRST_ROW - 2, column=2, value="Inputs").font = BOLD
    ws.cell(row=FIRST_ROW - 2, column=10, value="Calculated per row").font = BOLD

    for offset in range(ROW_COUNT):
        r = FIRST_ROW + offset
        ws.cell(row=r, column=1, value=offset + 1).alignment = Alignment(horizontal="center")
        if offset < len(example["rows"]):
            for col, value in enumerate(example["rows"][offset], start=2):
                ws.cell(row=r, column=col, value=value)
        guard = f'IF($B{r}="","",'
        ws[f"J{r}"] = f"={guard}D{r}+(ABS(E{r})-ABS(F{r}))/2)"
        ws[f"K{r}"] = f"={guard}(ABS(E{r})+ABS(F{r}))/2)"
        ws[f"L{r}"] = (
            f'={guard}IF(H{r}="Uniform",K{r}/SQRT(3),IF(H{r}="Triangular",K{r}/SQRT(6),K{r}/IF(I{r}>0,I{r},3))))'
        )
        ws[f"M{r}"] = f"={guard}G{r}*J{r})"
        ws[f"N{r}"] = f"={guard}ABS(G{r})*K{r})"
        ws[f"O{r}"] = f"={guard}(G{r}*L{r})^2)"
        ws[f"P{r}"] = f'={guard}IF(SUM($O${FIRST_ROW}:$O${LAST_ROW})=0,0,O{r}/SUM($O${FIRST_ROW}:$O${LAST_ROW})))'

    style_range(ws, f"B{FIRST_ROW}:I{LAST_ROW}", fill=INPUT_FILL)
    style_range(ws, f"A{FIRST_ROW}:A{LAST_ROW}", font=Font(color=MUTED))
    style_range(ws, f"J{FIRST_ROW}:P{LAST_ROW}", fill=CALC_FILL, font=Font(color=INK))
    style_range(ws, f"D{FIRST_ROW}:F{LAST_ROW}", number_format="0.000")
    style_range(ws, f"J{FIRST_ROW}:O{LAST_ROW}", number_format="0.0000")
    style_range(ws, f"P{FIRST_ROW}:P{LAST_ROW}", number_format="0.0%")

    dist_rule = DataValidation(type="list", formula1='"Normal,Uniform,Triangular"', allow_blank=True)
    dist_rule.error = "Choose Normal, Uniform or Triangular"
    dist_rule.prompt = "Normal: capable molded/machined. Uniform: flatness, float, clearance. Triangular: hand assembly."
    ws.add_data_validation(dist_rule)
    dist_rule.add(f"H{FIRST_ROW}:H{LAST_ROW}")
    k_rule = DataValidation(type="decimal", operator="between", formula1="0.5", formula2="10", allow_blank=True)
    k_rule.prompt = "Tolerance = k \u00d7 \u03c3. 3 \u2192 Cp 1.0, 4 \u2192 Cp 1.33"
    ws.add_data_validation(k_rule)
    k_rule.add(f"I{FIRST_ROW}:I{LAST_ROW}")
    units_rule = DataValidation(type="list", formula1='"mm,in"', allow_blank=False)
    ws.add_data_validation(units_rule)
    units_rule.add("C5")

    rng = lambda col: f"{col}{FIRST_ROW}:{col}{LAST_ROW}"  # noqa: E731
    define_name(wb, "GapMean", f"{sheet}!$S$6")
    define_name(wb, "GapSigma", f"{sheet}!$S$10")
    results = [
        (4, "Verdict", None),
        (5, "Nominal gap", f"=SUMPRODUCT({rng('G')},{rng('D')})"),
        (6, "Mean gap (incl. asymmetric tol.)", f"=SUM({rng('M')})"),
        (7, "Worst-case half band", f"=SUM({rng('N')})"),
        (8, "Worst case min", "=S6-S7"),
        (9, "Worst case max", "=S6+S7"),
        (10, "RSS \u03c3", f"=SQRT(SUM({rng('O')}))"),
        (11, "RSS \u22123\u03c3", "=S6-3*S10"),
        (12, "RSS +3\u03c3", "=S6+3*S10"),
        (13, "Cp", '=IF(OR(LSL="",USL="",S10=0),"",(USL-LSL)/(6*S10))'),
        (14, "Cpk", '=IF(AND(LSL="",USL=""),"",IF(S10=0,"",MIN(IF(LSL="",9E+99,S6-LSL),IF(USL="",9E+99,USL-S6))/(3*S10)))'),
        (15, "Predicted out of spec (ppm, normal)", '=IF(OR(AND(LSL="",USL=""),S10=0),"",(IF(LSL="",0,NORMSDIST((LSL-S6)/S10))+IF(USL="",0,1-NORMSDIST((USL-S6)/S10)))*1000000)'),
        (16, "Biggest lever (largest variance share)", f'=IF(SUM({rng("O")})=0,"",INDEX({rng("B")},MATCH(MAX({rng("P")}),{rng("P")},0)))'),
        (17, "Its variance share", f'=IF(SUM({rng("O")})=0,"",MAX({rng("P")}))'),
    ]
    ws["R3"] = "Results"
    ws["R3"].font = Font(bold=True, size=12, color=INK)
    for row, label, formula in results:
        ws.cell(row=row, column=18, value=label).font = BOLD
        if formula:
            ws.cell(row=row, column=19, value=formula)
    ws["S4"] = (
        '=IF(AND(LSL="",USL=""),"No spec limits",'
        'IF(AND(OR(LSL="",S8>=LSL),OR(USL="",S9<=USL)),"PASS \u2014 worst case",'
        'IF(AND(S14<>"",S14>=TargetCpk),"PASS (statistical) \u2014 WC exceeds spec","FAIL")))'
    )
    style_range(ws, "R4:S17", fill=RESULT_FILL)
    style_range(ws, "S5:S12", number_format="0.0000")
    style_range(ws, "S13:S14", number_format="0.00")
    style_range(ws, "S15", number_format="#,##0.00")
    style_range(ws, "S17", number_format="0.0%")
    ws["S4"].font = Font(bold=True, size=12)
    ws["S4"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[4].height = 32

    ws.conditional_formatting.add("S4", FormulaRule(formula=['LEFT($S$4,4)="FAIL"'], fill=PatternFill("solid", fgColor="FBD5DD"), font=Font(bold=True, color="9F1239")))
    ws.conditional_formatting.add("S4", FormulaRule(formula=['ISNUMBER(SEARCH("statistical",$S$4))'], fill=PatternFill("solid", fgColor="FDECC8"), font=Font(bold=True, color="92400E")))
    ws.conditional_formatting.add("S4", FormulaRule(formula=['LEFT($S$4,4)="PASS"'], fill=PatternFill("solid", fgColor="D3F4DF"), font=Font(bold=True, color="166534")))
    ws.conditional_formatting.add("S14", CellIsRule(operator="lessThan", formula=["TargetCpk"], font=Font(bold=True, color="BE123C")))
    ws.conditional_formatting.add("S14", CellIsRule(operator="greaterThanOrEqual", formula=["TargetCpk"], font=Font(bold=True, color="15803D")))
    ws.conditional_formatting.add(rng("P"), ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="num", end_value=0.5, end_color="7DD3FC"))

    ws["R19"] = "Tip: tighten the biggest lever first \u2014 halving a contributor that owns 35% of the variance cuts \u03c3 far more than halving one that owns 5%."
    ws["R19"].font = SUBTLE
    ws["R19"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("R19:S21")
    ws.freeze_panes = f"C{FIRST_ROW}"
    ws.print_options.gridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def build_lite(out_path):
    wb = Workbook()
    add_start_sheet(
        wb,
        "Lite",
        extra_lines=[
            ("Want more?", True),
            ("Stackup Pro adds Monte Carlo simulation with a histogram, a gasket compression checker, a DVT test tracker, an RCA / 8D issue log and a programme dashboard.", False),
            (f"Details: {APP_URL}#pro", False),
            ("", False),
        ],
    )
    stack = add_stack_sheet(wb)
    wb.active = wb.sheetnames.index(stack.title)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    target = Path(__file__).resolve().parent.parent / "web" / "downloads" / "Stackup-Lite.xlsx"
    print("wrote", build_lite(target))
